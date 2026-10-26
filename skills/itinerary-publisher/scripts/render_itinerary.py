#!/usr/bin/env python3
"""Render a structured travel itinerary to an editable DOCX."""

import json
import os
import re
import sys
import traceback
import uuid

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "itineraries")

DATE_FILL = "E0ECF8"
ACTIVITY_FILL = "FFFFE0"
HOTEL_FILL = "E0FFE0"
FLIGHT_FILL = "FFF0FF"
GROUND_TRANSPORT_FILL = "F0E6FF"
BORDER_COLOR = "666666"

# Defense in depth: the skill is instructed not to send pricing information,
# but the renderer also strips obvious pricing lines if one slips through.
_CURRENCY_RE = re.compile(
    r"(?:[$€£¥]\s?\d|\b(?:USD|EUR|GBP|JPY|CAD|AUD|CHF|CNY|RMB|HKD|AED|NPR|INR)\s*\d)",
    re.IGNORECASE,
)
_PRICE_LABEL_RE = re.compile(
    r"\b(?:price|fare|room rate|nightly rate|tax(?:es)?|fee(?:s)?|total|amount paid|balance due|cost)\b",
    re.IGNORECASE,
)


def _is_pricing_line(value):
    text = str(value or "").strip()
    if not text:
        return False
    return bool(_CURRENCY_RE.search(text) or (_PRICE_LABEL_RE.search(text) and re.search(r"\d", text)))


def _clean_lines(lines):
    return [str(line) for line in (lines or []) if not _is_pricing_line(line)]


def _clean_item(item):
    item = dict(item or {})
    if _is_pricing_line(item.get("title")):
        item["title"] = ""
    item["lines"] = _clean_lines(item.get("lines"))
    return item


def _is_empty_plan_placeholder(item):
    """Return True for filler text that represents an intentionally empty day."""
    title = str(item.get("title") or "").strip().lower()
    lines = [str(line).strip().lower() for line in (item.get("lines") or [])]
    text = " ".join([title, *lines]).strip()

    if not text:
        return True

    empty_phrases = (
        "no confirmed activities",
        "no confirmed activity",
        "no confirmed plans",
        "nothing scheduled",
        "nothing confirmed",
        "no scheduled activities",
        "no activities found",
    )
    if any(phrase in text for phrase in empty_phrases):
        return True

    # Also suppress bare filler titles when there is no substantive detail.
    if not lines and (title == "open day" or title == "free day" or
                      title.startswith("open day in ") or title.startswith("free day in ")):
        return True

    return False


def _is_transport_status_only(item):
    """Return True for transportation text that describes an ongoing state, not an action."""
    title = str(item.get("title") or "").strip().lower()
    lines = [str(line).strip().lower() for line in (item.get("lines") or [])]
    text = " ".join([title, *lines]).strip()

    if not text:
        return True

    # Suppress common continuation/status wording. These describe availability or
    # ongoing use rather than a transportation event that occurs on this day.
    status_phrases = (
        "rental car continues",
        "car rental continues",
        "rental car in use",
        "car in use",
        "continue with rental car",
        "continuing with rental car",
        "rental car available",
        "car available",
        "vehicle available",
    )
    if any(phrase in text for phrase in status_phrases):
        return True

    # Also catch vendor-specific wording such as "Avis car rental continues".
    if re.search(r"\b(?:car|vehicle)\s+rental\s+continues\b", text):
        return True

    return False


def _is_unscheduled_travel_instruction(item):
    """Return True for self-directed travel instructions that belong in plans."""
    title = str(item.get("title") or "").strip()
    lines = [str(line).strip() for line in (item.get("lines") or [])]
    text = " ".join([title, *lines]).strip().lower()

    if not text:
        return False

    # These are directions/instructions, not booked transportation events.
    return bool(re.match(
        r"^(?:drive|continue driving|continue to|head to|proceed to|make your way to)\b",
        text,
    ))


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER_COLOR)


def _set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _set_repeat_table_widths(table):
    widths = [Inches(1.685), Inches(3.877), Inches(1.938)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    return widths


def _style_paragraph(paragraph, centered=False, keep_with_next=False):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if centered:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_line(paragraph, text, bold=False):
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    return run


def _add_item(cell, item, blank_before=False):
    item = _clean_item(item)
    title = item.get("title") or ""
    lines = item.get("lines") or []
    if not title and not lines:
        return

    if blank_before and cell.paragraphs and cell.paragraphs[-1].text:
        p = cell.add_paragraph()
        _style_paragraph(p)

    p = cell.paragraphs[-1] if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text else cell.add_paragraph()
    _style_paragraph(p)
    if title:
        _add_line(p, title, bold=True)
    for line in lines:
        if p.text:
            p.add_run("\n")
        _add_line(p, line)



def _render_day(table, day):
    date_row = table.add_row()
    merged = date_row.cells[0].merge(date_row.cells[1]).merge(date_row.cells[2])
    merged.text = ""
    p = merged.paragraphs[0]
    _style_paragraph(p, centered=True, keep_with_next=True)
    run = p.add_run(day.get("date") or "")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    _set_cell_shading(merged, DATE_FILL)
    _set_cell_margins(merged, top=20, bottom=20, start=60, end=60)
    _set_cell_borders(merged)

    row = table.add_row()
    _set_row_cant_split(row)
    cells = row.cells
    for cell in cells:
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_margins(cell)
        _set_cell_borders(cell)
        _style_paragraph(cell.paragraphs[0])

    raw_transportation = [_clean_item(i) for i in (day.get("transportation") or [])]

    # If the model accidentally puts a self-directed travel instruction such as
    # "Drive to Bucharest" in transportation, move it to the middle column.
    moved_to_plans = [
        dict(i, kind=(i.get("kind") or "travel"))
        for i in raw_transportation
        if _is_unscheduled_travel_instruction(i)
    ]

    transportation = [
        i for i in raw_transportation
        if (i.get("title") or i.get("lines"))
        and not _is_transport_status_only(i)
        and not _is_unscheduled_travel_instruction(i)
    ]
    for i, item in enumerate(transportation):
        _add_item(cells[0], item, blank_before=(i > 0))
    if transportation:
        # Flights use the existing light pink/red fill. Any transportation-only
        # cell with no flights uses light purple. If a day mixes a flight with
        # another transportation mode, flight coloring takes precedence.
        modes = {str(item.get("mode") or "").strip().lower() for item in transportation}
        fill = FLIGHT_FILL if "flight" in modes else GROUND_TRANSPORT_FILL
        _set_cell_shading(cells[0], fill)

    plans = moved_to_plans + [_clean_item(i) for i in (day.get("plans") or [])]
    plans = [
        i for i in plans
        if (i.get("title") or i.get("lines")) and not _is_empty_plan_placeholder(i)
    ]
    for i, item in enumerate(plans):
        _add_item(cells[1], item, blank_before=(i > 0))
    # Pale yellow is reserved for cells containing only scheduled events.
    # Generic travel instructions, notes, and unscheduled activities remain unshaded.
    if plans and all((item.get("kind") == "scheduled_event") for item in plans):
        _set_cell_shading(cells[1], ACTIVITY_FILL)

    lodging = _clean_item(day.get("lodging") or {})
    if lodging.get("title") or lodging.get("lines"):
        title = lodging.get("title") or ""
        lines = lodging.get("lines") or []
        p = cells[2].paragraphs[0]
        if title:
            _add_line(p, title, bold=bool(lodging.get("full_details")))
        for line in lines:
            if p.text:
                p.add_run("\n")
            _add_line(p, line)
        if lodging.get("full_details"):
            _set_cell_shading(cells[2], HOTEL_FILL)


def build_itinerary(data):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    stem = f"travel-itinerary-{uuid.uuid4().hex}"
    filename = f"{stem}.docx"
    path = os.path.join(OUTPUT_DIR, filename)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)

    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for day in data.get("days") or []:
        _render_day(table, day)

    _set_repeat_table_widths(table)

    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.space_before = Pt(0)
        doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
        doc.paragraphs[-1].paragraph_format.line_spacing = 0.1

    doc.save(path)

    return path


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "expected one JSON argument"}))
        sys.exit(1)

    try:
        data = json.loads(sys.argv[1])
        if not isinstance(data, dict) or not isinstance(data.get("days"), list):
            raise ValueError("JSON must be an object containing a 'days' array")
        path = build_itinerary(data)
        print(json.dumps({"filename": os.path.basename(path)}))
    except Exception as e:
        print(json.dumps({
            "error": f"failed to render itinerary: {e}",
            "traceback": traceback.format_exc(),
        }))
        sys.exit(1)


if __name__ == "__main__":
    main()
